import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONTE = "Arial"
COR_CABECALHO = "1F3864"
CORES_STATUS = {
    "ENCONTRADO": "C6EFCE",      # verde
    "AMBIGUO": "FFEB9C",         # amarelo
    "NAO_ENCONTRADO": "FFC7CE",  # vermelho claro
}

LEGENDA_STATUS = [
    ("ENCONTRADO", "Rua localizada pelo nome exato ou com tipo compatível, e setor decidido "
                   "(setor único na rua, CEP ou número exato)."),
    ("AMBIGUO", "O setor é um palpite: há mais de um setor possível, ou a rua foi achada por "
                "correção de grafia ou com tipo diferente, ou o setor foi estimado pelo CEP ou pela "
                "localidade (FALLBACK_*). A coluna do setor traz o melhor palpite e "
                "setores_candidatos lista as alternativas."),
    ("NAO_ENCONTRADO", "Nenhuma rua compatível no município. O setor fica vazio."),
]

LEGENDA_COLUNAS = [
    ("setor_censitario_encontrado", "Código do setor no CNEFE (16 caracteres: 15 dígitos + letra). "
                                    "Vazio quando NAO_ENCONTRADO."),
    ("setor_censitario_15dig", "O mesmo setor sem a letra final."),
    ("status_vinculacao", "ENCONTRADO, AMBIGUO ou NAO_ENCONTRADO (veja acima)."),
    ("metodo_vinculacao", "Como o setor foi decidido, por exemplo RUA_EXATA+NUMERO_PROXIMO_CEP. "
                          "O prefixo FUZZY_ indica rua achada por correção de grafia; FALLBACK_ indica rua não "
                          "encontrada, com setor estimado pelo CEP ou pela localidade."),
    ("n_setores_candidatos", "Quantos setores diferentes existem na rua, entre os candidatos considerados."),
    ("setores_candidatos", "Até 5 setores mais frequentes, separados por | (apenas nos AMBIGUO)."),
]


def resumir(final):
    """Tabelas de contagem por status e por método (com percentual do total)."""
    total = len(final)
    contagem = final["status_vinculacao"].value_counts().reindex(list(CORES_STATUS), fill_value=0)
    por_status = contagem.rename_axis("status").reset_index(name="enderecos")
    por_status["percentual"] = por_status["enderecos"] / total
    por_metodo = (
        final["metodo_vinculacao"].value_counts().rename_axis("metodo").reset_index(name="enderecos")
    )
    por_metodo["percentual"] = por_metodo["enderecos"] / total
    return por_status, por_metodo


def _fonte_em_tudo(ws):
    for linha in ws.iter_rows():
        for celula in linha:
            celula.font = Font(name=FONTE, size=10)


def _cabecalho(ws, numero_linha):
    for celula in ws[numero_linha]:
        if celula.value is not None:
            celula.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", start_color=COR_CABECALHO)
            celula.alignment = Alignment(horizontal="center", vertical="center")


def _titulo(ws, numero_linha, texto):
    ws.cell(row=numero_linha, column=1, value=texto).font = Font(name=FONTE, size=11, bold=True)


def _larguras(ws, maximo=60):
    for coluna in ws.columns:
        tamanho = max(len(str(c.value)) for c in coluna if c.value is not None)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(tamanho * 1.15 + 2, maximo)


def _formatar_resultado(ws):
    _fonte_em_tudo(ws)
    _cabecalho(ws, 1)
    cabecalhos = [c.value for c in ws[1]]
    col_status = cabecalhos.index("status_vinculacao") + 1
    for linha in range(2, ws.max_row + 1):
        celula = ws.cell(row=linha, column=col_status)
        cor = CORES_STATUS.get(celula.value)
        if cor:
            celula.fill = PatternFill("solid", start_color=cor)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _larguras(ws)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _formatar_resumo(ws, n_status, inicio_metodo):
    _fonte_em_tudo(ws)
    _titulo(ws, 1, "Endereços por status")
    _cabecalho(ws, 2)
    _titulo(ws, inicio_metodo, "Endereços por método")
    _cabecalho(ws, inicio_metodo + 1)
    for linha in range(3, ws.max_row + 1):
        ws.cell(row=linha, column=3).number_format = "0.0%"
    for linha in range(3, 3 + n_status):
        cor = CORES_STATUS.get(ws.cell(row=linha, column=1).value)
        if cor:
            ws.cell(row=linha, column=1).fill = PatternFill("solid", start_color=cor)
    _larguras(ws)


def _formatar_legenda(ws, n_status, inicio_colunas):
    _fonte_em_tudo(ws)
    _titulo(ws, 1, "Significado dos status")
    _cabecalho(ws, 2)
    _titulo(ws, inicio_colunas, "Colunas acrescentadas à base de busca")
    _cabecalho(ws, inicio_colunas + 1)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    for linha in range(3, ws.max_row + 1):
        ws.cell(row=linha, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=linha, column=1).alignment = Alignment(vertical="top")


def salvar_excel(final, caminho):
    """Grava o resultado em Excel, com cores por status, resumo e legenda."""
    por_status, por_metodo = resumir(final)
    legenda_status = pd.DataFrame(LEGENDA_STATUS, columns=["status", "significado"])
    legenda_colunas = pd.DataFrame(LEGENDA_COLUNAS, columns=["coluna", "significado"])
    inicio_metodo = len(por_status) + 4
    inicio_colunas = len(legenda_status) + 4

    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        final.to_excel(escritor, sheet_name="resultado", index=False)
        por_status.to_excel(escritor, sheet_name="resumo", index=False, startrow=1)
        por_metodo.to_excel(escritor, sheet_name="resumo", index=False, startrow=inicio_metodo)
        legenda_status.to_excel(escritor, sheet_name="legenda", index=False, startrow=1)
        legenda_colunas.to_excel(escritor, sheet_name="legenda", index=False, startrow=inicio_colunas)

        livro = escritor.book
        _formatar_resultado(livro["resultado"])
        _formatar_resumo(livro["resumo"], len(por_status), inicio_metodo)
        _formatar_legenda(livro["legenda"], len(legenda_status), inicio_colunas)